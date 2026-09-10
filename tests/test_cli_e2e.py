# -*- coding: utf-8 -*-
"""端到端测试：生成 Excel/CSV → 生成 SQL → 用 SQLite 真实执行并校验数据。"""

import csv
import sqlite3
from datetime import datetime

from openpyxl import Workbook

from excel2sql.cli import main


def make_workbook(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "用户表"
    sheet.append(["姓名", "手机号", "余额", "注册日期", "备注"])
    sheet.append(["张三", "13800138000", 12.5, "2023-01-05", "含'引号'"])
    sheet.append(["李四", "007", 100, "2023-02-06", None])
    sheet.append([None, None, None, None, "只有备注"])
    order = workbook.create_sheet("订单")
    order.append(["订单编号", "金额", "创建时间"])
    order.append(["A001", 99.9, datetime(2023, 1, 5, 10, 30, 0)])
    order.append(["A002", 0.1, datetime(2023, 1, 6, 23, 59, 59)])
    workbook.save(path)


def run_cli(args):
    return main(args)


def read_sql(path):
    with open(path, "r", encoding="utf-8") as handle:
        return handle.read()


def test_sqlite_end_to_end(tmp_path):
    source = tmp_path / "数据.xlsx"
    output = tmp_path / "out.sql"
    make_workbook(str(source))

    assert run_cli([str(source), "-o", str(output), "--dialect", "sqlite", "-q"]) == 0
    sql = read_sql(output)

    assert 'CREATE TABLE IF NOT EXISTS "用户表"' in sql
    assert '"手机号" TEXT' in sql          # 智能文本识别：手机号按文本
    assert '"余额" NUMERIC' in sql
    assert '"注册日期" TEXT' in sql

    connection = sqlite3.connect(":memory:")
    connection.executescript(sql)
    rows = connection.execute(
        'SELECT 姓名, 手机号, 余额, 注册日期, 备注 FROM "用户表"'
    ).fetchall()
    assert len(rows) == 3
    assert ("张三", "13800138000", 12.5, "2023-01-05", "含'引号'") in rows
    assert ("李四", "007", 100, "2023-02-06", None) in rows

    order_rows = connection.execute(
        'SELECT 订单编号, 金额, 创建时间 FROM "订单" ORDER BY 订单编号'
    ).fetchall()
    assert len(order_rows) == 2
    assert order_rows[0][0] == "A001"
    assert order_rows[0][2] == "2023-01-05 10:30:00"
    connection.close()


def test_generic_sheet_name_falls_back_to_file_name(tmp_path):
    source = tmp_path / "成绩表.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(["姓名", "分数"])
    sheet.append(["王五", 90])
    workbook.save(str(source))

    output = tmp_path / "out.sql"
    assert run_cli([str(source), "-o", str(output), "--dialect", "sqlite", "-q"]) == 0
    sql = read_sql(output)
    assert 'CREATE TABLE IF NOT EXISTS "成绩表"' in sql
    assert '"Sheet1"' not in sql


def test_limit_and_add_id(tmp_path):
    source = tmp_path / "t.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "数据"
    sheet.append(["名称"])
    for index in range(5):
        sheet.append(["行%d" % index])
    workbook.save(str(source))

    output = tmp_path / "limit.sql"
    assert run_cli([str(source), "-o", str(output), "--dialect", "sqlite",
                    "--limit", "2", "--add-id", "-q"]) == 0
    sql = read_sql(output)
    assert '"id" INTEGER PRIMARY KEY AUTOINCREMENT' in sql

    connection = sqlite3.connect(":memory:")
    connection.executescript(sql)
    count = connection.execute('SELECT COUNT(*) FROM "数据"').fetchone()[0]
    ids = [row[0] for row in connection.execute('SELECT id FROM "数据" ORDER BY id')]
    assert count == 2
    assert ids == [1, 2]
    connection.close()


def test_all_text_option(tmp_path):
    source = tmp_path / "t.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["编号", "金额"])
    sheet.append(["A1", 10])
    workbook.save(str(source))

    output = tmp_path / "all_text.sql"
    assert run_cli([str(source), "-o", str(output), "--dialect", "sqlite",
                    "--all-text", "-q"]) == 0
    sql = read_sql(output)
    assert sql.count("TEXT") >= 2
    assert "NUMERIC" not in sql


def test_dry_run_writes_nothing(tmp_path):
    source = tmp_path / "t.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["a"])
    sheet.append([1])
    workbook.save(str(source))

    output = tmp_path / "should_not_exist.sql"
    assert run_cli([str(source), "-o", str(output), "--dry-run", "-q"]) == 0
    assert not output.exists()


def test_csv_input(tmp_path):
    source = tmp_path / "people.csv"
    with open(source, "w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["姓名", "工号", "工资"])
        writer.writerow(["赵六", "007", "8000.50"])

    output = tmp_path / "csv.sql"
    assert run_cli([str(source), "-o", str(output), "--dialect", "sqlite", "-q"]) == 0
    sql = read_sql(output)
    connection = sqlite3.connect(":memory:")
    connection.executescript(sql)
    row = connection.execute('SELECT 姓名, 工号, 工资 FROM "people"').fetchone()
    assert row == ("赵六", "007", 8000.5)
    connection.close()


def test_no_insert_only_ddl(tmp_path):
    source = tmp_path / "t.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["a"])
    sheet.append([1])
    workbook.save(str(source))

    output = tmp_path / "ddl.sql"
    assert run_cli([str(source), "-o", str(output), "--dialect", "sqlite",
                    "--no-insert", "-q"]) == 0
    sql = read_sql(output)
    assert "CREATE TABLE" in sql
    assert "INSERT INTO" not in sql


def test_sheet_filter(tmp_path):
    source = tmp_path / "t.xlsx"
    make_workbook(str(source))
    output = tmp_path / "one.sql"
    assert run_cli([str(source), "-o", str(output), "--dialect", "sqlite",
                    "--sheet", "订单", "-q"]) == 0
    sql = read_sql(output)
    assert '"订单"' in sql
    assert '"用户表"' not in sql


def test_table_name_and_prefix_suffix(tmp_path):
    source = tmp_path / "t.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "数据"
    sheet.append(["a"])
    sheet.append([1])
    workbook.save(str(source))

    output = tmp_path / "named.sql"
    assert run_cli([str(source), "-o", str(output), "--dialect", "sqlite",
                    "--table-name", "custom", "--prefix", "pre_", "--suffix", "_suf", "-q"]) == 0
    assert '"pre_custom_suf"' in read_sql(output)


def test_multiple_inputs_require_output(tmp_path):
    source = tmp_path / "t.xlsx"
    make_workbook(str(source))
    assert run_cli([str(source), str(source), "--dialect", "sqlite", "-q"]) == 2


def test_unknown_dialect_returns_error_code(tmp_path):
    source = tmp_path / "t.xlsx"
    make_workbook(str(source))
    assert run_cli([str(source), "--dialect", "oracle", "-q"]) == 2


def test_missing_input_returns_error_code(tmp_path):
    missing = tmp_path / "nope.xlsx"
    assert run_cli([str(missing), "-o", str(tmp_path / "o.sql"), "-q"]) == 1


def test_no_row_lost_when_sample_is_full(tmp_path):
    """回归：采样读满时不能丢掉那一行（历史 bug：10 万行只写出 99999 行）"""
    source = tmp_path / "rows.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "数据"
    sheet.append(["编号"])
    for index in range(1200):
        sheet.append(["N%04d" % index])
    workbook.save(str(source))

    output = tmp_path / "rows.sql"
    assert run_cli([str(source), "-o", str(output), "--dialect", "sqlite",
                    "--sample-size", "100", "-q"]) == 0
    connection = sqlite3.connect(":memory:")
    connection.executescript(read_sql(output))
    count = connection.execute('SELECT COUNT(*) FROM "数据"').fetchone()[0]
    assert count == 1200
    connection.close()


def test_mysql_dialect_defaults(tmp_path):
    source = tmp_path / "t.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "users"
    sheet.append(["id", "name"])
    sheet.append([1, "张三"])
    workbook.save(str(source))

    output = tmp_path / "mysql.sql"
    assert run_cli([str(source), "-o", str(output), "-q"]) == 0
    sql = read_sql(output)
    assert "SET NAMES utf8mb4;" in sql
    assert "ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;" in sql
    assert "`users`" in sql
    assert "SET FOREIGN_KEY_CHECKS = 0;" in sql
